from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models import AreaType, Dish, DishTag, FoodArea, FoodSection, FoodVendor, Tag, TagCategory
from app.schemas import ImportErrorItem, ImportResult, ImportRowResult


IMPORT_HEADERS = [
    "类型",
    "地点名称",
    "大概位置",
    "楼层名称",
    "楼层序号",
    "店铺/窗口名称",
    "店铺位置",
    "店铺图片URL",
    "菜品名称",
    "价格",
    "基础标签",
    "味道标签",
    "菜品简介",
    "图片URL",
    "是否启用",
    "是否供应",
]

REQUIRED_IMPORT_HEADERS = [header for header in IMPORT_HEADERS if header != "店铺图片URL"]

AREA_TYPE_ALIASES = {
    "餐厅": AreaType.CANTEEN,
    "校内餐厅": AreaType.CANTEEN,
    "CANTEEN": AreaType.CANTEEN,
    "小吃街": AreaType.SNACK_STREET,
    "外面的小吃街": AreaType.SNACK_STREET,
    "SNACK_STREET": AreaType.SNACK_STREET,
    "自定义方位": AreaType.CUSTOM_AREA,
    "自定义": AreaType.CUSTOM_AREA,
    "CUSTOM_AREA": AreaType.CUSTOM_AREA,
}


@dataclass
class ImportRow:
    row_number: int
    area_type: AreaType
    area_name: str
    location: str | None
    section_name: str | None
    section_level: int | None
    vendor_name: str
    vendor_location: str | None
    vendor_image_url: str | None
    dish_name: str
    price: Decimal | None
    basic_tags: list[str]
    taste_tags: list[str]
    dish_description: str | None
    image_url: str | None
    is_active: bool
    is_available: bool


def build_food_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "food-data"
    sheet.append(IMPORT_HEADERS)
    sheet.append(["餐厅", "文瀛餐厅", "南区", "一层", 1, "山西面馆", "东侧", "", "油泼刀削面", "12.50", "面", "辣", "香辣筋道", "", "是", "是"])
    sheet.append(["小吃街", "许西小吃街", "南门外", "", "", "东北饺子馆", "入口左侧", "", "招牌水饺", "15.00", "面", "鲜", "手工现包", "", "是", "是"])
    sheet.append(["自定义方位", "实验楼旁", "主楼西侧", "", "", "咖啡简餐", "靠近连廊", "", "鸡肉三明治", "11.00", "包子", "鲜", "适合早八", "", "是", "是"])
    for column in sheet.columns:
        header = column[0].value
        sheet.column_dimensions[column[0].column_letter].width = max(len(str(header)) + 6, 14)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def import_food_workbook(db: Session, content: bytes, dry_run: bool) -> ImportResult:
    rows, errors = parse_workbook(content)
    if errors:
        return ImportResult(
            dry_run=dry_run,
            created_count=0,
            updated_count=0,
            skipped_count=0,
            error_count=len(errors),
            errors=errors,
            rows=[ImportRowResult(row=error.row, status="error", message=error.message) for error in errors],
        )

    if dry_run:
        return simulate_import(db, rows)
    return apply_import(db, rows)


def parse_workbook(content: bytes) -> tuple[list[ImportRow], list[ImportErrorItem]]:
    try:
        workbook = load_workbook(BytesIO(content), data_only=True)
    except Exception:
        return [], [ImportErrorItem(row=0, message="Excel 文件无法读取，请上传 .xlsx 文件")]

    sheet = workbook.active
    headers = [normalize_cell(cell.value) for cell in sheet[1]]
    missing = [header for header in REQUIRED_IMPORT_HEADERS if header not in headers]
    if missing:
        return [], [ImportErrorItem(row=1, message=f"缺少列：{'、'.join(missing)}")]

    indexes = {header: headers.index(header) for header in IMPORT_HEADERS if header in headers}
    rows: list[ImportRow] = []
    errors: list[ImportErrorItem] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(is_blank(value) for value in values):
            continue
        data = {
            header: values[indexes[header]] if header in indexes and indexes[header] < len(values) else None
            for header in IMPORT_HEADERS
        }
        try:
            rows.append(parse_row(row_number, data))
        except ValueError as exc:
            errors.append(ImportErrorItem(row=row_number, message=str(exc)))
    return rows, errors


def parse_row(row_number: int, data: dict[str, object]) -> ImportRow:
    area_type_text = normalize_cell(data["类型"])
    area_type = AREA_TYPE_ALIASES.get(area_type_text)
    if area_type is None:
        raise ValueError("类型必须是 餐厅、小吃街、自定义方位 或对应英文枚举")

    area_name = required_text(data["地点名称"], "地点名称")
    vendor_name = required_text(data["店铺/窗口名称"], "店铺/窗口名称")
    dish_name = required_text(data["菜品名称"], "菜品名称")
    section_name = optional_text(data["楼层名称"])
    section_level = parse_level(data["楼层序号"])
    if area_type == AreaType.CANTEEN:
        section_level = section_level or 1
        section_name = section_name or f"{section_level}层"

    return ImportRow(
        row_number=row_number,
        area_type=area_type,
        area_name=area_name,
        location=optional_text(data["大概位置"]),
        section_name=section_name,
        section_level=section_level,
        vendor_name=vendor_name,
        vendor_location=optional_text(data["店铺位置"]),
        vendor_image_url=optional_text(data["店铺图片URL"]),
        dish_name=dish_name,
        price=parse_price(data["价格"]),
        basic_tags=split_tags(data["基础标签"]),
        taste_tags=split_tags(data["味道标签"]),
        dish_description=optional_text(data["菜品简介"]),
        image_url=optional_text(data["图片URL"]),
        is_active=parse_bool(data["是否启用"], default=True),
        is_available=parse_bool(data["是否供应"], default=True),
    )


def simulate_import(db: Session, rows: list[ImportRow]) -> ImportResult:
    created = 0
    updated = 0
    simulated_areas: set[tuple[AreaType, str]] = set()
    simulated_sections: set[tuple[AreaType, str, int]] = set()
    simulated_vendors: set[tuple[AreaType, str, int | None, str]] = set()
    simulated_dishes: set[tuple[AreaType, str, int | None, str, str]] = set()
    simulated_tags: set[tuple[TagCategory, str]] = set()
    row_results: list[ImportRowResult] = []

    for row in rows:
        row_created, row_updated = count_row(db, row, simulated_areas, simulated_sections, simulated_vendors, simulated_dishes, simulated_tags)
        created += row_created
        updated += row_updated
        status = "created" if row_created else "updated"
        row_results.append(ImportRowResult(row=row.row_number, status=status, message="预检通过"))

    return ImportResult(
        dry_run=True,
        created_count=created,
        updated_count=updated,
        skipped_count=0,
        error_count=0,
        errors=[],
        rows=row_results,
    )


def apply_import(db: Session, rows: list[ImportRow]) -> ImportResult:
    created = 0
    updated = 0
    row_results: list[ImportRowResult] = []
    for row in rows:
        area, area_created = get_or_create_area(db, row)
        section = None
        section_created = False
        if row.area_type == AreaType.CANTEEN:
            section, section_created = get_or_create_section(db, area, row)
        vendor, vendor_created = get_or_create_vendor(db, area, section, row)
        dish, dish_created = get_or_create_dish(db, vendor, row)
        tag_created_count = assign_import_tags(db, dish, row)
        created += sum([area_created, section_created, vendor_created, dish_created]) + tag_created_count
        updated += 0 if dish_created else 1
        row_results.append(
            ImportRowResult(row=row.row_number, status="created" if dish_created else "updated", message="导入成功")
        )
    db.commit()
    return ImportResult(
        dry_run=False,
        created_count=created,
        updated_count=updated,
        skipped_count=0,
        error_count=0,
        errors=[],
        rows=row_results,
    )


def count_row(
    db: Session,
    row: ImportRow,
    areas: set[tuple[AreaType, str]],
    sections: set[tuple[AreaType, str, int]],
    vendors: set[tuple[AreaType, str, int | None, str]],
    dishes: set[tuple[AreaType, str, int | None, str, str]],
    tags: set[tuple[TagCategory, str]],
) -> tuple[int, int]:
    created = 0
    updated = 0
    area_key = (row.area_type, row.area_name)
    area_exists = area_key in areas or db.query(FoodArea).filter_by(area_type=row.area_type, name=row.area_name).one_or_none()
    if area_exists:
        updated += 1
    else:
        created += 1
        areas.add(area_key)

    section_level = row.section_level if row.area_type == AreaType.CANTEEN else None
    if row.area_type == AreaType.CANTEEN and section_level is not None:
        section_key = (row.area_type, row.area_name, section_level)
        db_section_exists = (
            db.query(FoodSection).join(FoodSection.area).filter(FoodArea.area_type == row.area_type, FoodArea.name == row.area_name, FoodSection.level == section_level).one_or_none()
        )
        if section_key in sections or db_section_exists:
            updated += 1
        else:
            created += 1
            sections.add(section_key)

    vendor_key = (row.area_type, row.area_name, section_level, row.vendor_name)
    db_vendor_exists = (
        db.query(FoodVendor)
        .join(FoodVendor.area)
        .outerjoin(FoodVendor.section)
        .filter(FoodArea.area_type == row.area_type, FoodArea.name == row.area_name, FoodVendor.name == row.vendor_name)
        .one_or_none()
    )
    if vendor_key in vendors or db_vendor_exists:
        updated += 1
    else:
        created += 1
        vendors.add(vendor_key)

    dish_key = (*vendor_key, row.dish_name)
    db_dish_exists = (
        db.query(Dish)
        .join(Dish.vendor)
        .join(FoodVendor.area)
        .filter(FoodArea.area_type == row.area_type, FoodArea.name == row.area_name, FoodVendor.name == row.vendor_name, Dish.name == row.dish_name)
        .one_or_none()
    )
    if dish_key in dishes or db_dish_exists:
        updated += 1
    else:
        created += 1
        dishes.add(dish_key)

    for category, tag_name in row_tags(row):
        tag_key = (category, tag_name)
        tag_exists = tag_key in tags or db.query(Tag).filter_by(category=category, name=tag_name).one_or_none()
        if not tag_exists:
            created += 1
            tags.add(tag_key)
    return created, updated


def get_or_create_area(db: Session, row: ImportRow) -> tuple[FoodArea, bool]:
    area = db.query(FoodArea).filter_by(area_type=row.area_type, name=row.area_name).one_or_none()
    created = area is None
    if area is None:
        area = FoodArea(area_type=row.area_type, name=row.area_name)
        db.add(area)
        db.flush()
    area.location = row.location
    area.description = area.description or "批量导入"
    area.is_active = row.is_active
    return area, created


def get_or_create_section(db: Session, area: FoodArea, row: ImportRow) -> tuple[FoodSection, bool]:
    section = db.query(FoodSection).filter_by(area_id=area.id, level=row.section_level).one_or_none()
    created = section is None
    if section is None:
        level = row.section_level or 1
        section = FoodSection(area_id=area.id, level=level, name=row.section_name or f"{level}层")
        db.add(section)
        db.flush()
    section.name = row.section_name or f"{section.level}层"
    section.is_active = row.is_active
    return section, created


def get_or_create_vendor(db: Session, area: FoodArea, section: FoodSection | None, row: ImportRow) -> tuple[FoodVendor, bool]:
    vendor = db.query(FoodVendor).filter_by(area_id=area.id, section_id=section.id if section else None, name=row.vendor_name).one_or_none()
    created = vendor is None
    if vendor is None:
        vendor = FoodVendor(area_id=area.id, section_id=section.id if section else None, name=row.vendor_name)
        db.add(vendor)
        db.flush()
    vendor.description = vendor.description or "批量导入"
    vendor.location_hint = row.vendor_location
    vendor.image_url = row.vendor_image_url or vendor.image_url
    vendor.is_active = row.is_active
    return vendor, created


def get_or_create_dish(db: Session, vendor: FoodVendor, row: ImportRow) -> tuple[Dish, bool]:
    dish = db.query(Dish).filter_by(vendor_id=vendor.id, name=row.dish_name).one_or_none()
    created = dish is None
    if dish is None:
        dish = Dish(vendor_id=vendor.id, name=row.dish_name)
        db.add(dish)
        db.flush()
    dish.price = row.price
    dish.description = row.dish_description
    dish.image_url = row.image_url
    dish.is_available = row.is_available
    return dish, created


def assign_import_tags(db: Session, dish: Dish, row: ImportRow) -> int:
    created = 0
    tags: list[Tag] = []
    for category, tag_name in row_tags(row):
        tag = db.query(Tag).filter_by(category=category, name=tag_name).one_or_none()
        if tag is None:
            tag = Tag(category=category, name=tag_name, is_active=True)
            db.add(tag)
            db.flush()
            created += 1
        tags.append(tag)
    dish.tag_links = [DishTag(tag=tag) for tag in tags]
    return created


def row_tags(row: ImportRow) -> Iterable[tuple[TagCategory, str]]:
    for tag in row.basic_tags:
        yield TagCategory.BASIC, tag
    for tag in row.taste_tags:
        yield TagCategory.TASTE, tag


def normalize_cell(value: object) -> str:
    return str(value).strip() if value is not None else ""


def is_blank(value: object) -> bool:
    return normalize_cell(value) == ""


def required_text(value: object, label: str) -> str:
    text = normalize_cell(value)
    if not text:
        raise ValueError(f"{label}不能为空")
    return text


def optional_text(value: object) -> str | None:
    text = normalize_cell(value)
    return text or None


def parse_level(value: object) -> int | None:
    text = normalize_cell(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError as exc:
        raise ValueError("楼层序号必须是数字") from exc


def parse_price(value: object) -> Decimal | None:
    text = normalize_cell(value)
    if not text:
        return None
    try:
        price = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("价格必须是数字") from exc
    if price < 0:
        raise ValueError("价格不能小于 0")
    return price


def split_tags(value: object) -> list[str]:
    text = normalize_cell(value)
    if not text:
        return []
    normalized = text.replace("，", ",").replace("、", ",").replace("；", ",").replace(";", ",")
    return [tag.strip() for tag in normalized.split(",") if tag.strip()]


def parse_bool(value: object, default: bool) -> bool:
    text = normalize_cell(value)
    if not text:
        return default
    if text in {"是", "启用", "供应", "true", "TRUE", "1", "yes", "Y"}:
        return True
    if text in {"否", "停用", "停售", "false", "FALSE", "0", "no", "N"}:
        return False
    return default
